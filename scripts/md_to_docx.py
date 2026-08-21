#!/usr/bin/env python3
"""
Convert COMPLETE_INTEGRATED_MANUSCRIPT.md to an editable Word document.

Handles the subset of Markdown used by the manuscript: ATX headings, bold/italic
inline spans, bullet and numbered lists, blockquote callouts, fenced code blocks
(used for the fixed-width supplementary data tables), and horizontal rules.

Usage:  python3 scripts/md_to_docx.py [input.md] [output.docx]
"""

import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

IMAGE = re.compile(r'^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$')

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)')


def add_runs(paragraph, text):
    """Add text to a paragraph, honoring **bold**, *italic* and `code` spans."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        elif part.startswith('*') and part.endswith('*'):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def convert(md_path, docx_path):
    doc = Document()

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)

    lines = open(md_path, encoding='utf-8').read().split('\n')
    in_code = False
    code_buffer = []

    def flush_code():
        """Emit buffered fenced-code content as a monospace block."""
        if not code_buffer:
            return
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run('\n'.join(code_buffer))
        run.font.name = 'Consolas'
        run.font.size = Pt(7.5)
        code_buffer.clear()

    for line in lines:
        if line.startswith('```'):
            if in_code:
                flush_code()
            in_code = not in_code
            continue

        if in_code:
            code_buffer.append(line)
            continue

        stripped = line.strip()

        if not stripped:
            continue

        if stripped in ('---', '***', '___'):
            doc.add_paragraph('_' * 70).alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        img = IMAGE.match(stripped)
        if img:
            src = img.group('src')
            if os.path.exists(src):
                doc.add_picture(src, width=Inches(6.3))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                warn = doc.add_paragraph()
                warn.add_run(f'[missing image: {src}]').italic = True
            continue

        heading = re.match(r'^(#{1,6})\s+(.*)', stripped)
        if heading:
            level = len(heading.group(1))
            text = re.sub(r'\*\*|\*', '', heading.group(2))
            doc.add_heading(text, level=min(level, 4))
            continue

        # Blockquote callouts carry the simulated-data warnings; render them
        # visibly distinct so they cannot be mistaken for body text.
        if stripped.startswith('>'):
            text = stripped.lstrip('>').strip()
            if not text:
                continue
            text = re.sub(r'^#{1,6}\s+', '', text)
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(24)
            para.paragraph_format.space_after = Pt(2)
            add_runs(para, text)
            for run in para.runs:
                run.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
                run.font.size = Pt(10)
            continue

        bullet = re.match(r'^[-*+]\s+(.*)', stripped)
        if bullet:
            add_runs(doc.add_paragraph(style='List Bullet'), bullet.group(1))
            continue

        numbered = re.match(r'^\d+\.\s+(.*)', stripped)
        if numbered:
            add_runs(doc.add_paragraph(style='List Number'), numbered.group(1))
            continue

        add_runs(doc.add_paragraph(), stripped)

    flush_code()
    doc.save(docx_path)
    return docx_path


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else 'COMPLETE_INTEGRATED_MANUSCRIPT.md'
    dst = sys.argv[2] if len(sys.argv) > 2 else 'COMPLETE_INTEGRATED_MANUSCRIPT.docx'
    print(f'Wrote {convert(src, dst)}')
